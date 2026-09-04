# -*- coding: utf-8 -*-
"""Build a Word version of the paper from its LaTeX source.

Parses ../paper/main.tex (with \\input expanded) so the Word text is exactly the
paper's text, then lays it out IEEE-style: two columns, with the full-width
float (Fig. 1) and the wide results table placed in their own single-column
sections, as in the PDF.
"""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PAPER = '../paper'
SRC = open(f'{PAPER}/main.tex', encoding='utf-8').read()

# ---- expand \input so figures and table bodies are visible to the parser -----
def expand(t, depth=0):
    if depth > 3: return t
    def sub(m):
        p = os.path.join(PAPER, m.group(1))
        for cand in (p, p + '.tex'):
            if os.path.isfile(cand):
                return expand(open(cand, encoding='utf-8').read(), depth + 1)
        return ''
    return re.sub(r'\\input\{([^}]*)\}', sub, t)
SRC = expand(SRC)

# ---------------------------------------------------------------- inline markup
ACCENT = [(r"\'{c}", 'ć'), (r"\'{i}", 'í'), (r"\'{e}", 'é'), (r"\'{a}", 'á'),
          (r"\'{o}", 'ó'), (r"\'{u}", 'ú'), (r'\"{u}', 'ü'), (r'\"{o}', 'ö'),
          (r'\"{a}', 'ä'), (r'\~{n}', 'ñ'), (r'\~{a}', 'ã'), (r'\c{c}', 'ç'),
          (r'\c{C}', 'Ç'), (r'\v{c}', 'č'), (r'\v{s}', 'š'), (r'\^{o}', 'ô'),
          (r'\`{e}', 'è'), (r'\={o}', 'ō')]

GREEK = {r'\alpha':'α', r'\rho':'ρ', r'\Delta':'Δ', r'\times':'×', r'\approx':'≈',
         r'\geq':'≥', r'\leq':'≤', r'\in':'∈', r'\rightarrow':'→',
         r'\leftrightarrow':'↔', r'\star':'*', r'\pm':'±', r'\dots':'…', r'\cdot':'·'}
SUP = str.maketrans('0123456789-+', '⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺')
SUB = str.maketrans('0123456789', '₀₁₂₃₄₅₆₇₈₉')

def demath(m):
    s = m.group(1)
    s = re.sub(r'\\[,;!]', ' ', s)
    s = re.sub(r'\\math(?:rm|it|bf)\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\text(?:rm|it|bf)?\{([^{}]*)\}', r'\1', s)
    for k, v in GREEK.items(): s = s.replace(k, v)
    s = re.sub(r'\^\{([-+\d]+)\}', lambda x: x.group(1).translate(SUP), s)
    s = re.sub(r'\^(\d)', lambda x: x.group(1).translate(SUP), s)
    s = re.sub(r'_\{?([A-Za-z]{2,})\}?', r'_\1', s)
    s = re.sub(r'_\{?(\d)\}?', lambda x: x.group(1).translate(SUB), s)
    s = s.replace('{', '').replace('}', '').replace('\\', '')
    return re.sub(r'\s+', ' ', s).strip()

CITE, REF = {}, {}

def clean(t):
    t = re.sub(r'(?<!\\)%.*', '', t)
    t = re.sub(r'\\label\{[^}]*\}', '', t)
    t = re.sub(r'\\looseness\s*=\s*-?\d+', '', t)
    for a, b in ACCENT:                     # \'{c} etc. must go before ~ becomes a space
        t = t.replace(a, b)
    t = t.replace(r'\$', '\x1f')          # a literal dollar is not a math delimiter
    t = t.replace(r'\pex{}', 'PROMISE_exp').replace(r'\pex', 'PROMISE_exp')
    t = t.replace(r'\Fm{}', 'macro-F1').replace(r'\Fm', 'macro-F1')
    t = re.sub(r'\\cite\{([^}]*)\}', lambda m: '[' + ', '.join(
        CITE.get(k.strip(), '?') for k in m.group(1).split(',')) + ']', t)
    t = re.sub(r'\\eqref\{([^}]*)\}', lambda m: '(' + REF.get(m.group(1), '?') + ')', t)
    t = re.sub(r'\\ref\{([^}]*)\}', lambda m: REF.get(m.group(1), '?'), t)
    t = re.sub(r'\$([^$]*)\$', demath, t)
    t = re.sub(r'\\textbf\{((?:[^{}]|\{[^{}]*\})*)\}', '\x01' + r'\1' + '\x01', t)
    t = re.sub(r'\\(?:emph|textit)\{((?:[^{}]|\{[^{}]*\})*)\}', '\x02' + r'\1' + '\x02', t)
    t = re.sub(r'\\texttt\{([^{}]*)\}', '\x03' + r'\1' + '\x03', t)
    t = re.sub(r'\\textsc\{([^{}]*)\}', lambda m: m.group(1).upper(), t)
    t = re.sub(r'\\textsubscript\{([^{}]*)\}', '\x04' + r'\1' + '\x04', t)
    t = t.replace('---', '—').replace('--', '–')
    t = t.replace('``', '\u201c').replace("''", '\u201d')
    t = re.sub(r'\\[,;!]', ' ', t)
    t = re.sub(r'\\(?=\s)', '', t)      # "s.d.\ 0.198" and its line-broken twin
    for a, b in ((r'\%','%'), (r'\$','$'), (r'\&','&'), (r'\_','_'), (r'\#','#')):
        t = t.replace(a, b)
    t = t.replace('~', ' ')
    t = re.sub(r'\\[a-zA-Z]+\*?', '', t)
    t = t.replace('{', '').replace('}', '').replace('\x1f', '$')
    return re.sub(r'\s+', ' ', t).strip()

def add_runs(par, text, size=9.5, italic_all=False, bold_all=False):
    """Write text into a paragraph, honouring the \x01/\x02/\x03 markers."""
    tok = re.split(r'([\x01\x02\x03\x04])', text)
    b, i, m, sub = bold_all, italic_all, False, False
    for piece in tok:
        if piece == '\x01': b = not b; continue
        if piece == '\x02': i = not i; continue
        if piece == '\x03': m = not m; continue
        if piece == '\x04': sub = not sub; continue
        if not piece: continue
        r = par.add_run(piece)
        r.font.size = Pt(size); r.bold = b; r.italic = i
        r.font.subscript = sub
        r.font.name = 'Consolas' if m else 'Times New Roman'
        if m: r.font.size = Pt(size - 1)
    return par

# ---------------------------------------------------------------- numbering
for i, k in enumerate(re.findall(r'\\bibitem\{([^}]*)\}', SRC), 1):
    CITE[k] = str(i)
BIB = [(k, b.strip()) for k, b in re.findall(
    r'\\bibitem\{([^}]*)\}(.*?)(?=\\bibitem|\\end\{thebibliography\})', SRC, re.S)]

ROM = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
nf = nt = 0
for m in re.finditer(r'\\begin\{(figure\*?|table\*?)\}(.*?)\\end\{\1\}', SRC, re.S):
    lab = re.search(r'\\label\{([^}]*)\}', m.group(2))
    if not lab: continue
    if m.group(1).startswith('figure'):
        nf += 1; REF[lab.group(1)] = str(nf)
    else:
        nt += 1; REF[lab.group(1)] = ROM[nt]
for i, m in enumerate(re.finditer(r'\\section\{[^}]*\}(.*?)(?=\\section\{|\Z)', SRC, re.S), 1):
    body = m.group(1)
    for l in re.findall(r'\\label\{(sec:[^}]*)\}', body.split('\\subsection')[0]):
        REF[l] = ROM[i]
    for j, sb in enumerate(re.findall(r'\\subsection\{[^}]*\}(.*?)(?=\\subsection\{|\Z)', body, re.S)):
        for l in re.findall(r'\\label\{(sec:[^}]*)\}', sb):
            REF[l] = f'{ROM[i]}-{chr(65+j)}'
REF['eq:breakeven'] = '1'
print("figures:", nf, "| tables:", nt, "| refs:", len(BIB))
open('_refmap.txt','w').write(repr(REF))

# ---------------------------------------------------------------- docx helpers
def set_cols(section, n, space=340):
    cols = section._sectPr.xpath('./w:cols')[0]
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), str(space))
    cols.set(qn('w:equalWidth'), '1')

def new_section(doc, ncols):
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.68)
    s.top_margin = Inches(0.75); s.bottom_margin = Inches(0.9)
    set_cols(s, ncols)
    return s

def para(doc, text='', size=9.5, align=None, space_after=0, space_before=0,
         first_indent=None, italic=False, bold=False, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if first_indent is not None: pf.first_line_indent = Inches(first_indent)
    if align is not None: p.alignment = align
    if keep: pf.keep_with_next = True
    if text: add_runs(p, text, size, italic, bold)
    return p

def cell_borders(tbl, rows_top, rows_bottom):
    """booktabs look: rule above the header, under it, and at the foot."""
    for i, row in enumerate(tbl.rows):
        for c in row.cells:
            tcPr = c._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            for edge in ('top', 'bottom'):
                e = OxmlElement(f'w:{edge}')
                on = (edge == 'top' and i in rows_top) or (edge == 'bottom' and i in rows_bottom)
                e.set(qn('w:val'), 'single' if on else 'nil')
                e.set(qn('w:sz'), '8' if on else '0')
                e.set(qn('w:color'), '000000')
                b.append(e)
            tcPr.append(b)

def parse_tabular(block):
    # the column spec contains braces (@{}lccccc@{}), so match them balanced
    m = re.search(r'\\begin\{tabular\}\{(?:[^{}]|\{[^{}]*\})*\}(.*?)\\end\{tabular\}',
                  block, re.S)
    if not m: return None
    body = m.group(1)
    body = re.sub(r'\\(?:top|mid|bottom)rule', '', body)
    body = re.sub(r'\\cmidrule\([^)]*\)\{[^}]*\}', '', body)
    body = re.sub(r'\\cmidrule\{[^}]*\}', '', body)
    body = re.sub(r'\\addlinespace(\[[^\]]*\])?', '', body)
    rows = []
    for raw in re.split(r'\\\\', body):
        raw = raw.strip()
        if not raw: continue
        cells, span = [], []
        for c in re.split(r'(?<!\\)&', raw):
            mc = re.match(r'\s*\\multicolumn\{(\d+)\}\{[^}]*\}\{(.*)\}\s*$', c.strip(), re.S)
            if mc:
                cells.append(clean(mc.group(2))); span.append(int(mc.group(1)))
            else:
                cells.append(clean(c)); span.append(1)
        if any(x.strip() for x in cells):
            rows.append((cells, span))
    return rows

def add_table(doc, rows, size=8.0, header_rows=1, total_in=3.15, first_frac=0.34):
    ncol = max(sum(sp) for _, sp in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'dxa')
    w.set(qn('w:w'), str(int(total_in * 1440))); tblPr.append(w)
    first = Inches(total_in * first_frac)
    rest = Inches(total_in * (1 - first_frac) / max(1, ncol - 1))
    widths = [first] + [rest] * (ncol - 1)
    for j, col in enumerate(t.columns):
        col.width = widths[j]
    for cells, span in rows:
        r = t.add_row()
        ci = 0
        for text, sp in zip(cells, span):
            if ci >= ncol: break
            cell = r.cells[ci]
            if sp > 1 and ci + sp - 1 < ncol:
                cell = cell.merge(r.cells[ci + sp - 1])
            cell.width = widths[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, text, size)
            ci += sp
    cell_borders(t, rows_top={0}, rows_bottom={header_rows - 1, len(t.rows) - 1})
    return t

# ---------------------------------------------------------------- build
FIGPNG = {'fig:arch':'fig1_architecture.png', 'fig:inversion':'fig3_fig_inversion.png',
          'fig:prior':'fig4_fig_prior_shift.png', 'fig:cost':'fig5_fig_cost.png'}
FIGW   = {'fig:arch':7.0, 'fig:inversion':3.25, 'fig:prior':3.25, 'fig:cost':3.25}

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0

s0 = doc.sections[0]
s0.page_width, s0.page_height = Inches(8.5), Inches(11)
s0.left_margin = s0.right_margin = Inches(0.68)
s0.top_margin = Inches(0.75); s0.bottom_margin = Inches(0.9)
set_cols(s0, 1)

body = SRC.split(r'\maketitle', 1)[1].split(r'\end{document}')[0]
body = re.sub(r'\\begin\{abstract\}.*?\\end\{abstract\}', '', body, flags=re.S)
body = re.sub(r'\\begin\{IEEEkeywords\}.*?\\end\{IEEEkeywords\}', '', body, flags=re.S)
body = body.replace(r'\cmark', '✓').replace(r'\xmark', '✗')

# --- title block -------------------------------------------------------------
title = re.search(r'\\title\{(.*?)\}\s*\n\s*\n', SRC, re.S).group(1)
title = clean(title.replace(r'\\', ' '))
p = para(doc, title, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
authors = 'Fatma El-Zahraa Samir, Khaled T. Wassif, and Lamia AbouZeid'
para(doc, authors, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
para(doc, '//Faculty of Computers and Artificial Intelligence//, //Cairo University//, Giza, Egypt'
     .replace('//', '\x02'), size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
para(doc, 'fatmaelzahraasamirabdelfattah@gmail.com', size=10,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

new_section(doc, 2)

abstract = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', SRC, re.S).group(1)
p = para(doc, '', size=9, space_after=6)
r = p.add_run('Abstract—'); r.bold = True; r.italic = True
r.font.size = Pt(9); r.font.name = 'Times New Roman'
add_runs(p, clean(abstract), 9, bold_all=True)

kw = re.search(r'\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}', SRC, re.S).group(1)
p = para(doc, '', size=9, space_after=10)
r = p.add_run('Index Terms—'); r.bold = True; r.italic = True
r.font.size = Pt(9); r.font.name = 'Times New Roman'
add_runs(p, clean(kw), 9, bold_all=True)

# --- walk the body -----------------------------------------------------------
SPECIAL = re.compile(
    r'\\section\{(?P<sec>[^}]*)\}'
    r'|\\subsection\{(?P<sub>[^}]*)\}'
    r'|\\begin\{(?P<env>figure\*?|table\*?|equation|enumerate|thebibliography)\}'
    r'(?P<envbody>.*?)\\end\{(?P=env)\}', re.S)

secno = 0
cur_cols = 2
pos = 0
pending_indent = False

def flush_text(chunk):
    global pending_indent
    for blk in re.split(r'\n\s*\n', chunk):
        txt = clean(blk)
        if not txt: continue
        para(doc, txt, size=9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent=0.0 if not pending_indent else 0.16, space_after=0)
        pending_indent = True

def breakout(ncols):
    global cur_cols
    if cur_cols != ncols:
        new_section(doc, ncols); cur_cols = ncols

for m in SPECIAL.finditer(body):
    flush_text(body[pos:m.start()])
    pos = m.end()
    if m.group('sec'):
        breakout(2); secno += 1
        para(doc, f'{ROM[secno]}.  ' + clean(m.group('sec')).upper(), size=10,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=4, keep=True)
        pending_indent = False
        subno = 0
    elif m.group('sub'):
        breakout(2); subno += 1
        para(doc, f'{chr(64+subno)}.  ' + clean(m.group('sub')), size=9.5,
             italic=True, space_before=7, space_after=2, keep=True)
        pending_indent = False
    else:
        env, eb = m.group('env'), m.group('envbody')
        if env.startswith('figure'):
            lab = re.search(r'\\label\{([^}]*)\}', eb)
            lab = lab.group(1) if lab else ''
            cap = re.search(r'\\caption\{(.*?)\}\s*\n?\s*\\label', eb, re.S)
            cap = clean(cap.group(1)) if cap else ''
            wide = env.endswith('*')
            breakout(1 if wide else 2)
            if lab in FIGPNG and os.path.exists(FIGPNG[lab]):
                pp = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=3)
                pp.add_run().add_picture(FIGPNG[lab], width=Inches(FIGW[lab]))
            elif lab == 'fig:prompt':      # the prompt template, set as a boxed listing
                tf = re.search(r'\\raggedright(.*?)Answer:\}\}', eb, re.S)
                raw = (tf.group(1) + 'Answer:') if tf else ''
                raw = raw.replace(r'$\langle$', '\u27e8').replace(r'$\rangle$', '\u27e9')
                raw = raw.replace(r'$\dots$', '\u2026').replace(r'$k$', 'k')
                raw = re.sub(r'\\textcolor\{[^}]*\}\{', '{', raw)
                raw = re.sub(r'\\(?:rmfamily|itshape|upshape|ttfamily|scriptsize)', '', raw)
                raw = raw.replace(r'\hfill', '    ')
                raw = re.sub(r'\\\\\[[^\]]*\]', '\n', raw)
                raw = raw.replace('\\\\', '\n')
                box = doc.add_table(rows=1, cols=1); box.style = 'Table Grid'
                box.autofit = False
                box.columns[0].width = Inches(3.05)
                c = box.rows[0].cells[0]; c.width = Inches(3.05); c.text = ''
                first, blank = True, False
                for line in raw.split('\n'):
                    line = clean(line)
                    line = line.replace('"""' + ' ', '"""')
                    if not line:
                        if first or blank: continue
                        blank = True
                    else:
                        blank = False
                    pp = c.paragraphs[0] if first else c.add_paragraph()
                    first = False
                    pp.paragraph_format.space_after = Pt(0)
                    pp.paragraph_format.line_spacing = 1.0
                    r = pp.add_run(line); r.font.name = 'Consolas'; r.font.size = Pt(7)
            n = 'Fig. ' + REF.get(lab, '?')
            para(doc, '' + f'{n}.' + '' + f'  {cap}', size=8.5,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY if wide else WD_ALIGN_PARAGRAPH.LEFT,
                 space_after=8)
            pending_indent = False
        elif env.startswith('table'):
            lab = re.search(r'\\label\{([^}]*)\}', eb)
            lab = lab.group(1) if lab else ''
            cap = re.search(r'\\caption\{(.*?)\}\s*\n?\s*\\label', eb, re.S)
            cap = clean(cap.group(1)) if cap else ''
            wide = env.endswith('*')
            breakout(1 if wide else 2)
            n = 'Table ' + REF.get(lab, '?')
            para(doc, n.upper(), size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=8, space_after=1, keep=True)
            para(doc, cap, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3, keep=True)
            rows = parse_tabular(eb)
            if rows:
                hdr = 2 if lab in ('tab:coverage', 'tab:indomain', 'tab:transfer') else 1
                ff = {'tab:coverage':0.40, 'tab:corpus':0.26, 'tab:models':0.24,
                      'tab:indomain':0.24, 'tab:transfer':0.28}.get(lab, 0.32)
                add_table(doc, rows, size=7.5 if not wide else 8.5, header_rows=hdr,
                          total_in=7.0 if wide else 3.15, first_frac=ff)
            para(doc, '', size=6, space_after=6)
            pending_indent = False
        elif env == 'equation':
            breakout(2)
            para(doc, 'N* = C₀ / (c_alt − c_enc)          (1)', size=10,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_before=5, space_after=5)
            pending_indent = False
        elif env == 'enumerate':
            breakout(2)
            for i, it in enumerate(re.split(r'\\item', eb)[1:], 1):
                txt = clean(it)
                if not txt: continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.first_line_indent = Inches(-0.16)
                p.paragraph_format.space_after = Pt(3)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs(p, f'{i})  ' + txt, 9.5)
            pending_indent = False
        elif env == 'thebibliography':
            breakout(2)
            para(doc, 'REFERENCES', size=10, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=10, space_after=4, keep=True)
            for i, (k, txt) in enumerate(BIB, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.24)
                p.paragraph_format.first_line_indent = Inches(-0.24)
                p.paragraph_format.space_after = Pt(1.5)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs(p, f'[{i}]  ' + clean(txt), 8.0)
            pending_indent = False

flush_text(body[pos:])
doc.save('Paper.docx')
print('Paper.docx written')
